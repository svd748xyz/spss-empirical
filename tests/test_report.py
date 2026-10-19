"""Report layout follows a Chinese three-line table, without inventing statistics."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from spss_agent.interpret import frequencies
from spss_agent.oms import Cell
from spss_agent.report import build_report, fmt_p


def _border(cell, edge: str):
    borders = cell._tc.tcPr.find(qn("w:tcBorders"))
    element = borders.find(qn(f"w:{edge}"))
    return element.get(qn("w:val")), element.get(qn("w:sz"))


class ReportLayoutTest(unittest.TestCase):
    def test_p_keeps_leading_zero(self) -> None:
        self.assertEqual(fmt_p(0.046), "0.046")
        self.assertEqual(fmt_p(0.0004), "<0.001")

    def test_frequency_rows_keep_their_variable(self) -> None:
        cells = [
            Cell("FREQUENCIES", "Frequencies", ("男",), ("Frequency",), 44, "", "性别"),
            Cell("FREQUENCIES", "Frequencies", ("男",), ("Valid Percent",), 18.8, "", "性别"),
            Cell("FREQUENCIES", "Frequencies", ("是",), ("Frequency",), 54, "", "是否独生子女"),
            Cell("FREQUENCIES", "Frequencies", ("是",), ("Valid Percent",), 23.08, "", "是否独生子女"),
        ]
        rows = frequencies(cells)
        self.assertEqual(rows[0]["variable"], "性别")
        self.assertEqual(rows[0]["percent"], 18.8)
        self.assertEqual(rows[1]["variable"], "是否独生子女")
        self.assertEqual(rows[1]["percent"], 23.08)

    def test_tables_are_captioned_three_line_and_grouped(self) -> None:
        payload = {
            "title": "实习护生职业风险应对行为的影响因素研究",
            "methods": "有效问卷 234 份。",
            "license_note": "temporary",
            "software": r"IBM SPSS Statistics custom @ D:\spss\stats.exe",
            "labels": {"gender": "性别", "score": "职业风险应对行为", "exposure": "职业暴露预防行为", "total": "总量表"},
            "sections": [
                {
                    "kind": "reliability",
                    "heading": "（一）职业暴露预防行为的内部一致性",
                    "scale": "exposure",
                    "result": {"alpha": 0.843, "n_items": 12, "items": [{"name": "Q1", "citc": 0.43, "alpha_if_deleted": 0.836}]},
                },
                {
                    "kind": "reliability",
                    "scale": "total",
                    "result": {
                        "alpha": 0.934,
                        "n_items": 41,
                        "items": [{"name": f"Q{i}", "citc": 0.4, "alpha_if_deleted": 0.93} for i in range(41)],
                    },
                },
                {
                    "kind": "frequencies",
                    "vars": ["gender"],
                    "rows": [
                        {"labels": ("男",), "n": 44, "percent": 18.803, "variable": "性别"},
                        {"labels": ("女",), "n": 190, "percent": 81.197, "variable": "性别"},
                    ],
                },
                {
                    "kind": "ttest",
                    "group": "gender",
                    "results": [{
                        "outcome": "score",
                        "t": -3.98,
                        "df": 232,
                        "p": 0.0002,
                        "equal_var": True,
                        "groups": [
                            {"group": "男", "n": 44, "mean": 156.16, "sd": 18.69},
                            {"group": "女", "n": 190, "mean": 167.02, "sd": 15.72},
                        ],
                    }],
                },
                {
                    "kind": "regression",
                    "dependent": "score",
                    "models": [{
                        "r2": 0.220,
                        "adj_r2": 0.199,
                        "F": 10.654,
                        "p": 0.0001,
                        "coefficients": [
                            {"name": "(Constant)", "B": 157.466, "se": 9.960, "t": 15.810, "p": 0.0001},
                            {"name": "gender", "B": 6.632, "se": 2.618, "beta": 0.154, "t": 2.533, "p": 0.012, "vif": 1.08},
                        ],
                    }],
                },
            ],
            "discussion": ["横断问卷不能作因果结论。"],
        }
        with tempfile.TemporaryDirectory() as folder:
            dest = Path(folder) / "报告.docx"
            build_report(payload, dest)
            doc = Document(dest)
            text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
            self.assertNotIn("stats.exe", text)
            self.assertNotIn("temporary", text)
            self.assertNotIn("过程文件", text)
            self.assertNotIn("Q0", text)
            table_text = "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            self.assertIn("表1", table_text)
            self.assertIn("0.843", table_text)
            self.assertIn("0.934", table_text)
            self.assertIn("18.80", table_text)
            self.assertIn("156.16±18.69", table_text)
            self.assertIn("<0.001", table_text)
            self.assertIn("未标准化系数", table_text)
            self.assertIn("调整后的R²=0.199", table_text)
            self.assertAlmostEqual(doc.sections[0].page_width.cm, 21.0, places=1)
            caption = doc.tables[0].cell(0, 0)
            self.assertEqual(_border(caption, "bottom"), ("single", "12"))
            self.assertEqual(_border(caption, "left")[0], "nil")
            header = doc.tables[0].rows[1].cells[0]
            self.assertEqual(_border(header, "bottom"), ("single", "6"))


if __name__ == "__main__":
    unittest.main()
